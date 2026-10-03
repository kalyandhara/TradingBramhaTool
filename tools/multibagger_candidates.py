#!/usr/bin/env python3
"""
multibagger_candidates.py
Third screener angle: early-stage small/mid-cap candidates showing fresh
institutional conviction, screened for capital efficiency and balance-sheet
safety rather than already-obvious growth (see GARP screen in
promoter_profit_surprise.py for that angle).

Query (screener.in):
  Market Capitalization < 20000 AND
  OPM > 10 AND
  Return on capital employed > 15 AND
  Debt to equity < 1 AND
  Sales growth 3Years > 15 AND
  Change in promoter holding > -1 AND
  (Change in FII holding + Change in DII holding) > 1

Rationale (autopsy of the GARP screen's blind spots):
  - Market Cap < 20000 Cr: multibaggers come almost exclusively from
    small/mid-caps, where fresh FII/DII buying actually moves the float.
  - No TTM/5yr sales-growth requirement (unlike the GARP screen): an
    early-inflection multibagger often has a rough trailing year while
    institutions are already buying ahead of the turn. Only the 3yr
    growth base is required, so a weak recent quarter doesn't disqualify.
  - ROCE > 15 + OPM > 10: capital efficiency and margin quality, the
    actual multibagger engine — sales growth alone can mean low-margin,
    capital-hungry growth that never compounds shareholder value.
  - Debt to equity < 1: rules out balance-sheet risk, a common reason
    cheap "growth" stocks blow up instead of re-rating.
  - Change in promoter holding > -1: excludes the "smart money in,
    insiders out" distribution pattern, without requiring promoters to
    be actively buying (that's too strict a bar for most real cases).

UNMEASURED, and as of 2026-09-12 provably unmeasurable on this desk's free data
(B-73 Phase 0, `multibagger_cardinality.py`). Nothing below has been scored, and
that is not for want of trying: reconstructing this QUERY point-in-time leaves
**45 events against the desk's floor of 200**, and only 9 reach 252 days. Two
ceilings, one of them permanent --
  * the cached price/fundamental panel is LARGE-cap (median ~Rs 53,000 Cr; only
    27.8% of events fall under this screen's own Rs 20,000 Cr ceiling), and
  * yfinance serves FOUR annual statement columns, so with the 60-day filing lag a
    point-in-time "Sales growth 3Years" verdict first exists around mid-2025 --
    which caps formation dates at ~5 quarters however much ownership history is
    backfilled (and 20 quarters ARE available: that blocker was false, and
    clearing it changed nothing).
Read every column below as a CANDIDATE LIST, never as a measured edge.

The nearby question that COULD be powered was pre-registered and then MEASURED on
2026-09-13 (`multibagger_conviction_study.py`, 1,112 events): the
institutional-conviction core alone -- small cap + rising FII/DII + promoters not
selling -- returns **-2.07% @63d and -4.85% @126d against an equal-length hold of
the same name**, while its RAW returns (+7.0% / +13.4%) look excellent because the
market rose ~39%/yr underneath it. Split by era the sign FLIPS (-3.27 / +5.24 /
-8.19 @63d), so the honest reading is **no stable edge**, not a reliable negative.

That study is NOT a verdict on this screen and may not be cited as one: it drops
four of the seven legs below and measures months where "multibagger" means years.
What it does say is that the conviction leg this screen is built around carries no
stable 3-6 month edge on its own, and that any raw backtest of a long-only
small-cap screen in this window is measuring the market.

Promoter pledge % is not reliably scrapable from the public shareholding
table and is NOT screened here — manually sanity-check pledge levels
before acting on any candidate.

Price-to-52w-high is computed and shown as a column (not a hard filter)
since early-stage candidates often haven't broken out yet — momentum
confirmation is left to the reader's judgment.

Run manually:
  py -3 tools/multibagger_candidates.py SESSIONID
"""

import re
import sys
import time
from pathlib import Path
from datetime import datetime

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import screener_scraper as sc

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Run:  py -3 -m pip install openpyxl")

BASE_DIR = Path(__file__).resolve().parent.parent
import out_paths as _out

QUERY = (
    "Market Capitalization < 20000 AND\n"
    "OPM > 10 AND\n"
    "Return on capital employed > 15 AND\n"
    "Debt to equity < 1 AND\n"
    "Sales growth 3Years > 15 AND\n"
    "Change in promoter holding > -1 AND\n"
    "(Change in FII holding + Change in DII holding) > 1"
)

_GREEN = PatternFill("solid", fgColor="C6EFCE")
_AMBER = PatternFill("solid", fgColor="FFEB9C")
_RED = PatternFill("solid", fgColor="FFC7CE")


def fetch_holding_changes(session, slug: str) -> dict[str, float | None]:
    soup = sc.fetch_company_page(session, slug)
    if soup is None:
        return {"promoter": None, "fii": None, "dii": None, "price": None, "high": None}

    result = {}
    for key, label in (("promoter", "Promoter"), ("fii", "FII"), ("dii", "DII")):
        quarterly = sc._parse_quarterly_holding(soup, label)
        if len(quarterly) < 2:
            result[key] = None
            continue
        quarters_sorted = sorted(quarterly.keys(), key=lambda q: pd.Period(q, freq="Q"))
        latest, prev = quarters_sorted[-1], quarters_sorted[-2]
        result[key] = round(quarterly[latest] - quarterly[prev], 3)

    price, high = _parse_price_stats(soup)
    result["price"] = price
    result["high"] = high
    return result


def _parse_price_stats(soup) -> tuple[float | None, float | None]:
    """Self-sources current price and 52-week high from the top-ratios stat
    block, to compute how close the candidate is to breaking out."""
    ul = soup.find("ul", id="top-ratios")
    if ul is None:
        return None, None
    price, high = None, None
    for li in ul.find_all("li"):
        text = li.get_text(" ", strip=True)
        if text.startswith("Current Price"):
            m = re.search(r"([\d,]+\.?\d*)", text.replace("Current Price", ""))
            if m:
                price = float(m.group(1).replace(",", ""))
        elif text.startswith("High / Low"):
            nums = re.findall(r"([\d,]+\.?\d*)", text.replace("High / Low", ""))
            if nums:
                high = float(nums[0].replace(",", ""))
    return price, high


def build_dataset(sessionid: str, query: str) -> pd.DataFrame:
    session = sc._session(sessionid)
    rows = sc.fetch_all_results(session, query)
    if not rows:
        sys.exit("No results — check the query or whether the cookie is still valid.")

    log = sc.log
    log.info("Fetching promoter/FII/DII holding + price stats for %d companies...", len(rows))
    for i, row in enumerate(rows, 1):
        changes = fetch_holding_changes(session, row["_slug"])
        row["Change in promoter holding (computed)"] = changes["promoter"]
        row["Change in FII holding (computed)"] = changes["fii"]
        row["Change in DII holding (computed)"] = changes["dii"]
        row["Current Price (computed)"] = changes["price"]
        row["52W High (computed)"] = changes["high"]
        time.sleep(1.5)
        if i % 25 == 0:
            log.info("  lookup %d/%d", i, len(rows))

    df = pd.DataFrame(rows).drop(columns=["_slug"], errors="ignore")
    for col in df.columns:
        if col == "Name":
            continue
        df[col] = pd.to_numeric(df[col].astype(str).str.replace(",", ""), errors="coerce")

    fii_col = "Change in FII holding (computed)"
    dii_col = "Change in DII holding (computed)"
    promoter_col = "Change in promoter holding (computed)"
    price_col = "Current Price (computed)"
    high_col = "52W High (computed)"

    df["Combined FII+DII Change"] = df[[fii_col, dii_col]].sum(axis=1, skipna=True)
    df["% of 52W High"] = (df[price_col] / df[high_col] * 100).round(1)
    df = df.sort_values("Combined FII+DII Change", ascending=False).reset_index(drop=True)

    for key, col in (("promoter", promoter_col), ("FII", fii_col), ("DII", dii_col)):
        missing = df[col].isna().sum()
        if missing:
            log.warning("%d/%d stocks missing %s holding data", missing, len(df), key)

    df.attrs["fii_col"] = fii_col
    df.attrs["dii_col"] = dii_col
    df.attrs["promoter_col"] = promoter_col
    return df


def write_excel_sheet(df: pd.DataFrame, xlsx_path: Path) -> None:
    if xlsx_path.exists():
        wb = openpyxl.load_workbook(xlsx_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    sheet_name = "Multibagger Candidates"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    ws["A1"] = "Early-stage small/mid-cap candidates: capital efficiency + fresh institutional conviction"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Report generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(italic=True, color="5F5E5A")
    ws["A3"] = QUERY.replace("\n", " ")
    ws["A3"].font = Font(size=9, color="5F5E5A")
    ws["A4"] = ("Holding-change and price/52W-high columns marked (computed) are self-sourced "
                "from each company's page. Promoter pledge % is NOT screened — sanity-check "
                "manually before acting. % of 52W High is informational only, not a filter — "
                "early candidates often haven't broken out yet.")
    ws["A4"].font = Font(size=9, italic=True, color="5F5E5A")
    header_row = 6

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = Font(bold=True)

    for r, row in enumerate(df.itertuples(index=False), header_row + 1):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)

    col_map = {name: idx + 1 for idx, name in enumerate(df.columns)}
    combined_idx = col_map["Combined FII+DII Change"]
    high_pct_idx = col_map.get("% of 52W High")
    for r in range(header_row + 1, header_row + 1 + len(df)):
        val = ws.cell(row=r, column=combined_idx).value or 0
        ws.cell(row=r, column=combined_idx).fill = _GREEN if val >= 3 else (_AMBER if val >= 1 else _RED)
        if high_pct_idx:
            hp = ws.cell(row=r, column=high_pct_idx).value or 0
            ws.cell(row=r, column=high_pct_idx).fill = _GREEN if hp >= 80 else (_AMBER if hp >= 50 else _RED)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    for col_idx, col_name in enumerate(df.columns, 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(col_name)), *(len(str(v)) for v in df[col_name].head(50)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 35)

    wb.save(xlsx_path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: py -3 tools/multibagger_candidates.py SESSIONID")

    sessionid = sys.argv[1]
    df = build_dataset(sessionid, QUERY)

    date_str = datetime.today().strftime("%Y%m%d")
    csv_path = _out.out_path(f"multibagger_candidates_{date_str}.csv")
    df.to_csv(csv_path, index=False)
    print(f"Saved -> {csv_path}  ({len(df)} rows)")

    xlsx_path = _out.out_path(f"sector_rotation_{date_str}.xlsx")
    write_excel_sheet(df, xlsx_path)
    print(f"Added 'Multibagger Candidates' sheet -> {xlsx_path}")
