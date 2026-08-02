#!/usr/bin/env python3
"""
sector_rotation_trail.py
Historical version of sector_rotation.py: builds a multi-quarter trail per
sector instead of a single current snapshot, using:
  - real quarterly FII% history per stock (screener.in's Shareholding
    Pattern table, captured by screener_scraper.py's --query run; see
    FII_<quarter> columns in its output CSV)
  - real trailing-1Y sector-index returns as of each quarter-end date
    (yfinance daily price history)

For each sector, produces a chronological sequence of
(quarter, fii_rotation_score, index_trailing_1y_return) points — this is
the data a rotation-quadrant chart needs to draw a historical tail and an
implied heading arrow (last point minus second-to-last point), instead of
a single dot.

Caveat: cap-weighting uses each stock's CURRENT market cap as a static
weight across all historical quarters (we don't have historical market
caps from this data source) — a reasonable approximation for sector
ranking, not a true historical cap-weighted index.

Run manually:
  py -3 tools/sector_rotation_trail.py screener_export_<date>.csv
"""

import sys
import json
import time
from pathlib import Path
from datetime import datetime

import pandas as pd

try:
    import yfinance as yf
except ImportError:
    sys.exit("Run:  py -3 -m pip install yfinance")

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.collections import LineCollection
except ImportError:
    sys.exit("Run:  py -3 -m pip install matplotlib")

try:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
except ImportError:
    sys.exit("Run:  py -3 -m pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sector_rotation import SECTOR_INDEX_MAP, index_for_industry

BASE_DIR = Path(__file__).resolve().parent.parent


def load_wide_export(path: Path) -> tuple[pd.DataFrame, list[str]]:
    df = pd.read_csv(path)
    quarter_cols = [c for c in df.columns if c.startswith("FII_")]
    quarter_cols.sort(key=lambda c: pd.Period(c.replace("FII_", ""), freq="Q"))
    return df, quarter_cols


def quarter_label_to_date(label: str) -> pd.Timestamp:
    return pd.Period(label, freq="Q").end_time.normalize()


def build_sector_fii_trail(df: pd.DataFrame, quarter_cols: list[str]) -> pd.DataFrame:
    rows = []
    for i in range(1, len(quarter_cols)):
        prev_col, cur_col = quarter_cols[i - 1], quarter_cols[i]
        sub = df[["Industry", "Mar Cap", prev_col, cur_col]].dropna()
        if sub.empty:
            continue
        sub = sub.copy()
        sub["delta"] = sub[cur_col] - sub[prev_col]
        sub["weighted_delta"] = sub["delta"] * sub["Mar Cap"]

        grouped = sub.groupby("Industry").agg(
            stock_count=("delta", "size"),
            total_mar_cap=("Mar Cap", "sum"),
            weighted_delta_sum=("weighted_delta", "sum"),
        )
        grouped["fii_rotation_score"] = grouped["weighted_delta_sum"] / grouped["total_mar_cap"]
        grouped["quarter"] = cur_col.replace("FII_", "")
        rows.append(grouped.reset_index())

    return pd.concat(rows, ignore_index=True) if rows else pd.DataFrame()


def fetch_index_history(ticker: str, start: pd.Timestamp, end: pd.Timestamp) -> pd.Series | None:
    try:
        df = yf.download(ticker, start=start.strftime("%Y-%m-%d"),
                          end=end.strftime("%Y-%m-%d"), interval="1d",
                          progress=False, auto_adjust=True)
    except Exception:
        return None
    if df is None or df.empty:
        return None
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = [c[0] for c in df.columns]
    return df["Close"] if "Close" in df.columns else df.iloc[:, 0]


def trailing_1y_return_asof(close: pd.Series, asof: pd.Timestamp) -> float | None:
    close = close.copy()
    close.index = pd.to_datetime(close.index)
    end_price = close.asof(asof)
    start_price = close.asof(asof - pd.Timedelta(days=365))
    if pd.isna(end_price) or pd.isna(start_price) or start_price == 0:
        return None
    return round((end_price / start_price - 1) * 100, 2)


def build_index_trail(sectors: list[str], quarter_labels: list[str]) -> dict[str, dict[str, float | None]]:
    quarter_dates = {q: quarter_label_to_date(q) for q in quarter_labels}
    earliest = min(quarter_dates.values()) - pd.Timedelta(days=400)
    latest = max(quarter_dates.values()) + pd.Timedelta(days=5)

    ticker_cache: dict[str, pd.Series | None] = {}
    result: dict[str, dict[str, float | None]] = {}

    for sector in sectors:
        # Same two-vocabulary lookup the snapshot uses — the orbit chart must
        # resolve an NSE-sourced frame as readily as a screener.in export.
        ticker = index_for_industry(sector)
        if ticker is None:
            result[sector] = {q: None for q in quarter_labels}
            continue
        if ticker not in ticker_cache:
            ticker_cache[ticker] = fetch_index_history(ticker, earliest, latest)
            time.sleep(1.5)
        close = ticker_cache[ticker]
        if close is None:
            result[sector] = {q: None for q in quarter_labels}
            continue
        result[sector] = {q: trailing_1y_return_asof(close, d) for q, d in quarter_dates.items()}

    return result


_PALETTE = ["#534AB7", "#0F6E56", "#993556", "#3B6D11", "#185FA5",
            "#993C1D", "#854F0B", "#D85A30", "#7F77DD", "#D4537E"]
_HOT = (217, 75, 50)


def _lerp(a: float, b: float, t: float) -> float:
    return a + (b - a) * t


def _segment_color(base_rgb: tuple[int, int, int], t: float) -> tuple[float, float, float, float]:
    pale = [_lerp(255, c, 0.4) for c in base_rgb]
    hot = [_lerp(c, h, 0.75) for c, h in zip(base_rgb, _HOT)]
    r, g, b = (_lerp(pale[i], hot[i], t) / 255 for i in range(3))
    alpha = _lerp(0.45, 1.0, t)
    return (r, g, b, alpha)


def render_orbit_chart(combined: list[dict], top_n: int, out_png: Path) -> None:
    ranked = sorted(combined, key=lambda d: d["trail"][-1]["fii_score"], reverse=True)
    ranked = [d for d in ranked if d["trail"][-1]["stock_count"] >= 3][:top_n]

    fig, ax = plt.subplots(figsize=(11, 7), facecolor="#FBF4E9")
    ax.set_facecolor("#FBF4E9")

    all_idx = [p["index_return"] for d in ranked for p in d["trail"] if p["index_return"] is not None]
    all_fii = [p["fii_score"] for d in ranked for p in d["trail"]]
    x_min, x_max = (min(all_idx) - 10, max(all_idx) + 10) if all_idx else (-25, 75)
    y_min, y_max = (min(all_fii) - 1, max(all_fii) + 1) if all_fii else (-4, 6)

    ax.axhspan(0, y_max, x_min, 0, color="#FAC775", alpha=0.18)
    ax.axhspan(0, y_max, 0, x_max, color="#97C459", alpha=0.16)
    ax.axhspan(y_min, 0, x_min, 0, color="#F09595", alpha=0.16)
    ax.axhspan(y_min, 0, 0, x_max, color="#F0997B", alpha=0.16)
    ax.axhline(0, color="#444441", lw=0.8, ls="--", alpha=0.4)
    ax.axvline(0, color="#444441", lw=0.8, ls="--", alpha=0.4)

    ax.text(x_max * 0.97, y_max * 0.95, "Leading", color="#27500A", fontsize=10, ha="right", fontweight="bold")
    ax.text(x_min * 0.95 if x_min < 0 else x_min + 1, y_max * 0.95, "Improving", color="#633806", fontsize=10, fontweight="bold")
    ax.text(x_max * 0.97, y_min * 0.9, "Weakening", color="#712B13", fontsize=10, ha="right", fontweight="bold")
    ax.text(x_min * 0.95 if x_min < 0 else x_min + 1, y_min * 0.9, "Lagging", color="#791F1F", fontsize=10, fontweight="bold")

    for i, d in enumerate(ranked):
        base_hex = _PALETTE[i % len(_PALETTE)]
        base_rgb = tuple(int(base_hex[j:j+2], 16) for j in (1, 3, 5))
        pts = [(p["index_return"], p["fii_score"]) for p in d["trail"] if p["index_return"] is not None]
        if len(pts) < 2:
            continue

        segments = [[pts[i], pts[i + 1]] for i in range(len(pts) - 1)]
        n_seg = len(segments)
        colors = [_segment_color(base_rgb, i / max(n_seg - 1, 1)) for i in range(n_seg)]
        widths = [_lerp(1.2, 2.6, i / max(n_seg - 1, 1)) for i in range(n_seg)]
        lc = LineCollection(segments, colors=colors, linewidths=widths, linestyles=":")
        ax.add_collection(lc)

        x0, y0 = pts[-2]
        x1, y1 = pts[-1]
        hot_color = tuple(_lerp(c, h, 0.75) / 255 for c, h in zip(base_rgb, _HOT))
        ax.annotate("", xy=(x1, y1), xytext=(x0, y0),
                    arrowprops=dict(arrowstyle="-|>", color=hot_color, lw=0, mutation_scale=18))
        ax.text(x1 + (x1 - x0) * 0.08, y1 + (y1 - y0) * 0.08, d["sector"],
                fontsize=8.5, color="#2C2C2A", fontweight="medium")

    ax.set_xlim(x_min, x_max)
    ax.set_ylim(y_min, y_max)
    ax.set_xlabel("Index 1Y return (%)", color="#444441", fontsize=10)
    ax.set_ylabel("FII rotation score", color="#444441", fontsize=10)
    ax.tick_params(colors="#5F5E5A", labelsize=9)
    for spine in ax.spines.values():
        spine.set_color("#D3D1C7")

    fig.tight_layout()
    fig.savefig(out_png, dpi=160, facecolor="#FBF4E9")
    plt.close(fig)


_SHORT_NAMES = {
    "Business Process Outsourcing (BPO)/ Knowledge Process Outsourcing (KPO)": "BPO/KPO",
    "Depositories, Clearing Houses and Other Intermediaries": "Depositories",
    "Healthcare Research, Analytics & Technology": "Healthcare R&A",
    "Non Banking Financial Company (NBFC)": "NBFC",
    "Telecom -  Equipment & Accessories": "Telecom Equip.",
    "Telecom - Cellular & Fixed line services": "Telecom Cellular",
    "Telecom - Infrastructure": "Telecom Infra",
    "Residential, Commercial Projects": "Real Estate",
    "Compressors, Pumps & Diesel Engines": "Compressors/Pumps",
    "Auto Components & Equipments": "Auto Components",
    "Gems, Jewellery And Watches": "Gems & Jewellery",
    "Furniture, Home Furnishing": "Furniture",
    "Plastic Products - Industrial": "Plastic (Ind.)",
    "Financial Technology (Fintech)": "Fintech",
    "Computers - Software & Consulting": "IT Software",
    "Healthcare Service Provider": "Healthcare Svc.",
    "Internet & Catalogue Retail": "Internet Retail",
    "Passenger Cars & Utility Vehicles": "Passenger Cars",
    "Oil Exploration & Production": "Oil E&P",
    "Medical Equipment & Supplies": "Medical Equip.",
    "Asset Management Company": "AMC",
    "Microfinance Institutions": "Microfinance",
    "Pesticides & Agrochemicals": "Agrochemicals",
}


def _short_name(name: str) -> str:
    if name in _SHORT_NAMES:
        return _SHORT_NAMES[name]
    return name[:16].strip() + "…" if len(name) > 18 else name


def _classify_bias_label(fii: float, idx: float) -> str:
    inflow, strong = fii > 0, idx > 0
    if inflow and strong:
        return "Leading · confirmed bullish rotation"
    if inflow and not strong:
        return "Improving · early accumulation"
    if not inflow and strong:
        return "Weakening · distribution into strength"
    return "Lagging · confirmed bearish rotation"


_ORBIT_HTML_TEMPLATE = """<!doctype html>
<html><head><meta charset="utf-8"><title>Sector Rotation Orbit Map</title>
<style>
  body { font-family: -apple-system, Segoe UI, Arial, sans-serif; background: #FBF4E9; margin: 0; padding: 20px; }
  .layout { display: grid; grid-template-columns: minmax(0,1fr) 280px; gap: 16px; max-width: 1400px; margin: 0 auto; }
  .chart-box { position: relative; width: 100%; height: 640px; background: #FBF4E9; border-radius: 12px; border: 1px solid #D3D1C7; padding: 12px; box-sizing: border-box; }
  .panel { background: #F1EFE8; border-radius: 12px; padding: 10px; max-height: 640px; overflow-y: auto; box-sizing: border-box; }
  .panel button { flex: 1; font-size: 12px; padding: 6px; cursor: pointer; border-radius: 6px; border: 1px solid #B4B2A9; background: white; }
  .sector-row { display: flex; align-items: flex-start; gap: 6px; padding: 6px 4px; font-size: 12px; cursor: pointer; border-radius: 6px; }
  .sector-row:hover { background: rgba(0,0,0,0.04); }
  .sector-row .bias { color: #5F5E5A; font-size: 10.5px; }
  h1 { font-size: 18px; font-weight: 500; max-width: 1400px; margin: 0 auto 4px; }
  .meta { font-size: 12px; color: #5F5E5A; max-width: 1400px; margin: 0 auto 16px; }
</style></head>
<body>
<h1>Sector rotation orbit map</h1>
<p class="meta">Report generated: __GENERATED_AT__ &middot; __SECTOR_COUNT__ sectors &middot; quarterly path __FIRST_QUARTER__ to __LAST_QUARTER__</p>
<div class="layout">
  <div class="chart-box"><canvas id="orbitChart"></canvas></div>
  <div class="panel">
    <div style="display:flex; gap:6px; margin-bottom:8px;">
      <button id="selectAllBtn">Select all</button>
      <button id="deselectAllBtn">Deselect all</button>
    </div>
    <p style="font-size:11px; font-weight:600; color:#5F5E5A; margin:0 0 6px; text-transform:uppercase; letter-spacing:0.03em;">All __SECTOR_COUNT__ sectors, ranked</p>
    <div id="sectorList"></div>
  </div>
</div>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.js"></script>
<script>
const rawSectors = __SECTORS_JSON__;
const SHORT_NAMES = __SHORT_NAMES_JSON__;

function shortName(name) {
  if (SHORT_NAMES[name]) return SHORT_NAMES[name];
  return name.length > 18 ? name.slice(0, 16).trim() + "\\u2026" : name;
}
function classifyBias(fii, idx) {
  const inflow = fii > 0, strong = idx > 0;
  if (inflow && strong) return "Leading \\u00b7 confirmed bullish rotation";
  if (inflow && !strong) return "Improving \\u00b7 early accumulation";
  if (!inflow && strong) return "Weakening \\u00b7 distribution into strength";
  return "Lagging \\u00b7 confirmed bearish rotation";
}

const HOT = [217, 75, 50];
function hexToRgb(hex) { return [parseInt(hex.slice(1,3),16), parseInt(hex.slice(3,5),16), parseInt(hex.slice(5,7),16)]; }
function lerp(a, b, t) { return a + (b - a) * t; }
function hotColorFor(baseRgb) {
  return "rgb(" + Math.round(lerp(baseRgb[0],HOT[0],0.75)) + "," + Math.round(lerp(baseRgb[1],HOT[1],0.75)) + "," + Math.round(lerp(baseRgb[2],HOT[2],0.75)) + ")";
}

const RANK_DARK = [35, 30, 80];
const RANK_LIGHT = [222, 220, 240];
const allSectors = rawSectors
  .map(s => ({ ...s, latestFii: s.points[s.points.length - 1].fii, latestIdx: s.points[s.points.length - 1].idx }))
  .sort((a, b) => b.latestFii - a.latestFii);
allSectors.forEach((s, i) => {
  const t = allSectors.length > 1 ? i / (allSectors.length - 1) : 0;
  const baseRgb = [0, 1, 2].map(k => Math.round(lerp(RANK_DARK[k], RANK_LIGHT[k], t)));
  s.baseColor = "#" + baseRgb.map(c => c.toString(16).padStart(2, "0")).join("");
  s.color = hotColorFor(baseRgb);
  s.bias = classifyBias(s.latestFii, s.latestIdx);
});

const defaultVisible = new Set(allSectors.slice(0, 5).map(s => s.name));
const xMin = -50, xMax = 100, yMin = -10, yMax = 20;

function segmentColor(baseRgb, t) {
  const pale = [lerp(255, baseRgb[0], 0.4), lerp(255, baseRgb[1], 0.4), lerp(255, baseRgb[2], 0.4)];
  const hot = [lerp(baseRgb[0], HOT[0], 0.75), lerp(baseRgb[1], HOT[1], 0.75), lerp(baseRgb[2], HOT[2], 0.75)];
  const r = Math.round(lerp(pale[0], hot[0], t)), g = Math.round(lerp(pale[1], hot[1], t)), b = Math.round(lerp(pale[2], hot[2], t));
  return "rgba(" + r + "," + g + "," + b + "," + lerp(0.45, 1, t).toFixed(2) + ")";
}

const quadrantPlugin = {
  id: "quadrantShading",
  beforeDraw(chart) {
    const { ctx, chartArea, scales } = chart;
    const x0 = scales.x.getPixelForValue(0);
    const y0 = scales.y.getPixelForValue(0);
    ctx.save();
    ctx.fillStyle = "rgba(151,196,89,0.16)"; ctx.fillRect(x0, chartArea.top, chartArea.right - x0, y0 - chartArea.top);
    ctx.fillStyle = "rgba(250,199,117,0.2)"; ctx.fillRect(chartArea.left, chartArea.top, x0 - chartArea.left, y0 - chartArea.top);
    ctx.fillStyle = "rgba(240,153,123,0.18)"; ctx.fillRect(x0, y0, chartArea.right - x0, chartArea.bottom - y0);
    ctx.fillStyle = "rgba(240,149,149,0.18)"; ctx.fillRect(chartArea.left, y0, x0 - chartArea.left, chartArea.bottom - y0);
    ctx.strokeStyle = "rgba(68,68,65,0.3)"; ctx.setLineDash([4, 4]);
    ctx.beginPath(); ctx.moveTo(x0, chartArea.top); ctx.lineTo(x0, chartArea.bottom); ctx.moveTo(chartArea.left, y0); ctx.lineTo(chartArea.right, y0); ctx.stroke();
    ctx.setLineDash([]);
    ctx.font = "500 12px sans-serif";
    ctx.fillStyle = "#27500A"; ctx.fillText("Leading", chartArea.right - 60, chartArea.top + 16);
    ctx.fillStyle = "#633806"; ctx.fillText("Improving", chartArea.left + 8, chartArea.top + 16);
    ctx.fillStyle = "#712B13"; ctx.fillText("Weakening", chartArea.right - 72, chartArea.bottom - 10);
    ctx.fillStyle = "#791F1F"; ctx.fillText("Lagging", chartArea.left + 8, chartArea.bottom - 10);
    ctx.restore();
  }
};

const arrowLabelPlugin = {
  id: "arrowLabel",
  afterDatasetsDraw(chart) {
    const { ctx } = chart;
    chart.data.datasets.forEach((ds) => {
      if (ds.hidden) return;
      const meta = chart.getDatasetMeta(chart.data.datasets.indexOf(ds));
      const pts = meta.data;
      if (!pts || pts.length < 2) return;
      const a = pts[pts.length - 2], b = pts[pts.length - 1];
      const angle = Math.atan2(b.y - a.y, b.x - a.x);
      const len = 10;
      ctx.save();
      ctx.fillStyle = ds._hotColor;
      ctx.translate(b.x, b.y); ctx.rotate(angle);
      ctx.beginPath(); ctx.moveTo(len, 0); ctx.lineTo(-len * 0.6, len * 0.55); ctx.lineTo(-len * 0.6, -len * 0.55); ctx.closePath(); ctx.fill();
      ctx.restore();
      ctx.save();
      ctx.font = "500 10px sans-serif";
      ctx.fillStyle = "#2C2C2A";
      ctx.fillText(ds._shortLabel, b.x + Math.cos(angle) * 14 + 3, b.y + Math.sin(angle) * 14 + 3);
      ctx.restore();
    });
  }
};

function makeDataset(s) {
  const baseRgb = hexToRgb(s.baseColor);
  const total = s.points.length - 1;
  return {
    label: s.name,
    data: s.points.map((p, i) => ({ x: p.idx, y: p.fii, q: p.q, count: p.count, name: s.name, isLast: i === s.points.length - 1 })),
    borderColor: s.color, backgroundColor: s.color,
    showLine: true, fill: false, borderWidth: 2,
    pointRadius: (ctx) => ctx.raw && ctx.raw.isLast ? 0 : 2,
    pointHoverRadius: (ctx) => ctx.raw && ctx.raw.isLast ? 7 : 4,
    tension: 0.45, cubicInterpolationMode: "monotone",
    segment: {
      borderDash: () => [2, 3],
      borderColor: (ctx) => segmentColor(baseRgb, total > 0 ? ctx.p0DataIndex / (total - 1 || 1) : 0),
      borderWidth: (ctx) => lerp(1.3, 2.6, total > 0 ? ctx.p0DataIndex / (total - 1 || 1) : 0)
    },
    hidden: !defaultVisible.has(s.name),
    _hotColor: s.color,
    _shortLabel: shortName(s.name)
  };
}

const chart = new Chart(document.getElementById("orbitChart"), {
  type: "scatter",
  data: { datasets: allSectors.map(makeDataset) },
  options: {
    responsive: true, maintainAspectRatio: false, parsing: false, layout: { padding: 16 },
    scales: {
      x: { min: xMin, max: xMax, title: { display: true, text: "Index 1Y return (%)", color: "#444441", font: { size: 12 } }, grid: { color: "rgba(68,68,65,0.08)" }, ticks: { color: "#5F5E5A" } },
      y: { min: yMin, max: yMax, title: { display: true, text: "FII rotation score", color: "#444441", font: { size: 12 } }, grid: { color: "rgba(68,68,65,0.08)" }, ticks: { color: "#5F5E5A" } }
    },
    plugins: {
      legend: { display: false },
      tooltip: {
        backgroundColor: "#2C2C2A", titleFont: { size: 13, weight: "500" }, bodyFont: { size: 12 }, padding: 10,
        filter: (item) => !chart.data.datasets[item.datasetIndex].hidden,
        callbacks: {
          title: (items) => items[0].raw.name + " - " + items[0].raw.q,
          label: (item) => {
            const d = item.raw;
            return ["FII rotation score: " + d.y.toFixed(2), "Index 1Y return: " + d.x.toFixed(2) + "%", d.count + " stocks in this sector that quarter"];
          }
        }
      }
    }
  },
  plugins: [quadrantPlugin, arrowLabelPlugin]
});

const listEl = document.getElementById("sectorList");
const checkboxes = [];
allSectors.forEach((s, i) => {
  const stockCount = s.points[s.points.length - 1].count;
  const row = document.createElement("label");
  row.className = "sector-row";
  row.title = s.name + ": " + stockCount + " stock" + (stockCount === 1 ? "" : "s") + " in this sector";
  row.innerHTML = '<input type="checkbox" ' + (defaultVisible.has(s.name) ? "checked" : "") + ' style="margin:2px 0 0 0;accent-color:' + s.color + ';flex-shrink:0;">' +
    '<span style="width:8px;height:8px;border-radius:2px;background:' + s.color + ';flex-shrink:0;margin-top:4px;"></span>' +
    '<span style="display:flex;flex-direction:column;line-height:1.3;">' +
      '<span>' + (i+1) + '. ' + s.name + ' (' + stockCount + ' stock' + (stockCount === 1 ? "" : "s") + ')</span>' +
      '<span class="bias">' + s.bias + '</span>' +
    '</span>';
  const checkbox = row.querySelector("input");
  checkboxes.push(checkbox);
  checkbox.addEventListener("change", () => { chart.data.datasets[i].hidden = !checkbox.checked; chart.update(); });
  listEl.appendChild(row);
});

document.getElementById("selectAllBtn").addEventListener("click", () => {
  checkboxes.forEach((cb, i) => { cb.checked = true; chart.data.datasets[i].hidden = false; });
  chart.update();
});
document.getElementById("deselectAllBtn").addEventListener("click", () => {
  checkboxes.forEach((cb, i) => { cb.checked = false; chart.data.datasets[i].hidden = true; });
  chart.update();
});
</script>
</body></html>
"""


def render_orbit_html(combined: list[dict], out_html: Path) -> None:
    sectors_for_html = []
    for d in combined:
        pts = [p for p in d["trail"] if p["index_return"] is not None]
        if len(pts) < 2:
            continue
        sectors_for_html.append({
            "name": d["sector"],
            "points": [{"q": p["quarter"], "fii": round(p["fii_score"], 2),
                        "idx": round(p["index_return"], 2), "count": p["stock_count"]} for p in pts],
        })

    all_quarters = sorted({p["q"] for s in sectors_for_html for p in s["points"]},
                           key=lambda q: pd.Period(q, freq="Q"))

    html = (_ORBIT_HTML_TEMPLATE
            .replace("__GENERATED_AT__", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            .replace("__SECTOR_COUNT__", str(len(sectors_for_html)))
            .replace("__FIRST_QUARTER__", all_quarters[0] if all_quarters else "n/a")
            .replace("__LAST_QUARTER__", all_quarters[-1] if all_quarters else "n/a")
            .replace("__SECTORS_JSON__", json.dumps(sectors_for_html, separators=(",", ":")))
            .replace("__SHORT_NAMES_JSON__", json.dumps(_SHORT_NAMES)))

    out_html.write_text(html, encoding="utf-8")


def write_excel_sheet(combined: list[dict], png_path: Path, html_path: Path, xlsx_path: Path, top_n: int) -> None:
    if xlsx_path.exists():
        wb = openpyxl.load_workbook(xlsx_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    sheet_name = "Rotation Map"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    ws["A1"] = "Sector rotation orbit map"
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    ws["A2"] = f"Report generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = openpyxl.styles.Font(italic=True, color="5F5E5A")
    ws["A3"] = f"Top {top_n} sectors by latest quarter FII rotation score (min. 3 stocks), quarterly path Sep 2023 onward"
    ws["A3"].font = openpyxl.styles.Font(size=10, color="5F5E5A")
    ws["A4"] = f"Full interactive version (all sectors, select all/deselect all, rotation bias): {html_path.name}"
    ws["A4"].font = openpyxl.styles.Font(size=10, color="185FA5")

    img = XLImage(str(png_path))
    img.width, img.height = 880, 560
    ws.add_image(img, "A6")

    wb.save(xlsx_path)


def run(export_path: Path) -> None:
    df, quarter_cols = load_wide_export(export_path)
    if len(quarter_cols) < 2:
        sys.exit(f"Need at least 2 FII_<quarter> columns to compute deltas, found {len(quarter_cols)}. "
                  "Re-run screener_scraper.py to capture quarterly history first.")

    fii_trail = build_sector_fii_trail(df, quarter_cols)
    if fii_trail.empty:
        sys.exit("No sector FII trail could be built — check the input data.")

    sectors = sorted(fii_trail["Industry"].unique())
    quarter_labels = [c.replace("FII_", "") for c in quarter_cols[1:]]

    print(f"Fetching index history for sector trails ({len(sectors)} sectors)...")
    index_trail = build_index_trail(sectors, quarter_labels)

    combined = []
    for sector in sectors:
        sector_rows = fii_trail[fii_trail["Industry"] == sector].set_index("quarter")
        points = []
        for q in quarter_labels:
            if q not in sector_rows.index:
                continue
            fii_score = float(sector_rows.loc[q, "fii_rotation_score"])
            idx_return = index_trail.get(sector, {}).get(q)
            stock_count = int(sector_rows.loc[q, "stock_count"])
            points.append({"quarter": q, "fii_score": round(fii_score, 3),
                            "index_return": idx_return, "stock_count": stock_count})
        if points:
            combined.append({"sector": sector, "trail": points})

    date_str = datetime.today().strftime("%Y%m%d")
    out_path = BASE_DIR / f"sector_rotation_trail_{date_str}.json"
    out_path.write_text(json.dumps(combined, indent=2))
    print(f"Saved -> {out_path}  ({len(combined)} sectors)")

    png_path = BASE_DIR / f"sector_rotation_orbit_{date_str}.png"
    render_orbit_chart(combined, top_n=10, out_png=png_path)
    print(f"Saved -> {png_path}")

    html_path = BASE_DIR / f"sector_rotation_orbit_{date_str}.html"
    render_orbit_html(combined, html_path)
    print(f"Saved -> {html_path}  (interactive: all sectors, select all/deselect all, rotation bias)")

    xlsx_path = BASE_DIR / f"sector_rotation_{date_str}.xlsx"
    write_excel_sheet(combined, png_path, html_path, xlsx_path, top_n=10)
    print(f"Added 'Rotation Map' sheet -> {xlsx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: py -3 tools/sector_rotation_trail.py <screener_export_with_quarters.csv>")
    run(Path(sys.argv[1]))
